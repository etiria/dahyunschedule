#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate bias-methodology guide as .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()
def set_base_font(name="Malgun Gothic", size=10.5):
    st = DOC.styles["Normal"]; st.font.name=name; st.font.size=Pt(size)
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), name)
set_base_font()
def _k(run, name="Malgun Gothic"):
    run.font.name=name; r=run._element.get_or_add_rPr(); rf=r.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); r.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), name)
def H(t,l=1):
    h=DOC.add_heading(level=l); _k(h.add_run(t)); return h
def P(t="",bold=False,italic=False,size=None,color=None,mono=False):
    p=DOC.add_paragraph(); r=p.add_run(t); _k(r, "Consolas" if mono else "Malgun Gothic")
    r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor(*color)
    if mono: r.font.size=Pt(8.5)
    return p
def BUL(t):
    p=DOC.add_paragraph(style="List Bullet"); _k(p.add_run(t)); return p
def NUM(t):
    p=DOC.add_paragraph(style="List Number"); _k(p.add_run(t)); return p
def CODE(lines):
    for ln in lines.split("\n"):
        p=DOC.add_paragraph(); r=p.add_run(ln if ln else " "); _k(r,"Consolas")
        r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x1f,0x49,0x7d)
        p.paragraph_format.left_indent=Cm(0.5); p.paragraph_format.space_after=Pt(0)
    DOC.add_paragraph()
def TABLE(headers, rows, widths=None):
    t=DOC.add_table(rows=1, cols=len(headers)); t.style="Light Grid Accent 1"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""; r=c.paragraphs[0].add_run(h); _k(r); r.bold=True; r.font.size=Pt(9)
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=""; r=cs[i].paragraphs[0].add_run(str(v)); _k(r); r.font.size=Pt(9)
    if widths:
        for i,w in enumerate(widths):
            for row in t.rows: row.cells[i].width=Cm(w)
    DOC.add_paragraph(); return t

# Title
tt=DOC.add_paragraph(); tt.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=tt.add_run("후향적 코호트에서 비뚤림(Bias) 회피 방법론\n— 상세 해설 및 예시"); _k(r); r.bold=True; r.font.size=Pt(16)
s=DOC.add_paragraph(); s.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=s.add_run("부속문서 · 서울성모병원 KT·H. pylori 1단계 연구 · 2026-06-10"); _k(r); r.italic=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x66,0x66,0x66)
DOC.add_paragraph()

H("0. 전체 지도 (어떤 비뚤림이, 왜 생기나)",1)
TABLE(["비뚤림","본 연구에서의 발생 지점","주 해결법"],
 [["Immortal time","KT(time-zero)~제균 시작 사이 '불멸기간'","시간의존 노출·landmark·CCW"],
  ["Selection","HP 검사받은 사람만 노출 분류; 추적손실; 유병사례","restriction+IPSW·다중대치·new-user"],
  ["Confounding by indication","제균·내시경은 증상/궤양자에 시행","PS(IPTW/매칭)·active comparator·E-value"],
  ["Detection/surveillance","이식·제균군이 내시경 더 자주→더 발견","내시경빈도 보정·결과 정의 강화"],
  ["Competing risk","사망·이식신장소실이 암/출혈을 '막음'","CIF·Fine-Gray"],
  ["Misclassification","코드 기반 노출/결과 오분류","검증된 정의·민감도·음성대조"],
  ["Lead-time/length","암 감시로 조기·완만한 암 더 포착","발생률 비교·임상유의 사건 제한"]],
 widths=[3.3,6.2,5.0])
P("공통 원칙: 분석 전 DAG(인과그래프)로 교란·collider·매개를 구분하고, target trial(가상 RCT)을 정의한 뒤 모방한다.", italic=True, size=9.5)

H("1. Immortal Time Bias (불멸시간 비뚤림)",1)
H("1.1 정의·기전",2)
P("추적 시작(time-zero)부터 노출 확정 시점까지 환자는 반드시 생존/무사건이어야 한다. 이 '불멸기간'을 노출군 person-time에 넣으면 노출군이 인위적으로 유리해진다. 본 연구: time-zero=KT일. '이식 후 제균자'를 제균군으로 보면 제균 시점까지 출혈/암·사망이 없어야 하므로 가짜 보호효과.")
H("1.2 잘못된 분석 예시 (숫자)",2)
BUL("진실은 제균 효과 없음(HR=1.0). 제균은 평균 이식 후 180일에 시행.")
BUL("Naive(ever/never): 제균군 첫 180일은 '출혈 없이 제균까지 도달한 사람'만 → 출혈이 면제 → 가짜 HR≈0.6 (순수 인공물).")
P("직관: '약 받을 때까지 살아남은 시간'을 약의 공으로 돌리는 오류(오스카 수상자 장수 착시).", italic=True, size=9.5)
H("1.3 해결 A — 시간의존 노출",2)
CODE("""library(survival)
dt <- tmerge(base, base, id=id, tstop=fu_time)
dt <- tmerge(dt, erad_dates, id=id, eradicated = tdc(erad_day))
coxph(Surv(tstart, tstop, bleed) ~ eradicated + age + sex, data=dt)
# -> 가상예 HR ≈ 1.0 복원""")
H("1.4 해결 B — Landmark 분석",2)
P("고정 시점(예 180일)까지의 제균 여부로 군을 나누고 그 이후부터 추적. 이전 사건·사망 제외. 단점: 정보손실·시점민감 → 6/12개월 민감도.")
CODE("""land <- subset(cohort, fu_time > 180)
land$grp <- ifelse(land$erad_day <= 180, "erad","no")
coxph(Surv(180, fu_time, bleed) ~ grp + age + sex, data=land)""")
H("1.5 해결 C — Target Trial Emulation + Clone-Censor-Weight",2)
NUM("Clone: 각 환자를 제균전략/비제균전략에 복제 배정(time-zero에 양 전략과 양립).")
NUM("Censor: 추적 중 자기 전략과 어긋나는 순간 인위적 중도절단.")
NUM("Weight: 인위적 중도절단 선택을 IPCW(역확률중도절단가중)로 보정.")
P("장점: 불멸시간·시점선택을 구조적으로 제거, RCT에 대응되는 per-protocol 효과.", italic=True, size=9.5)
H("1.6 주의",2)
P("감염상태 비교(HP− vs HP+ 미제균)는 baseline 노출이라 불멸시간 무관. 시간의존·CCW는 제균(G1 vs G2)에만 적용.")

H("2. Selection Bias (선택 비뚤림)",1)
H("2.1 정의·종류",2)
BUL("Ascertainment: HP 검사받은 사람만 HP+/− 분류 → '검사받음'이 selection.")
BUL("Loss to follow-up: 결과와 연관된 탈락. / Prevalent user: 생존자 편향.")
H("2.2 본 연구 핵심: HP 검사 selection",2)
P("HP 검사는 증상/내시경자에게 시행 → 검사군이 일반 KT와 다름. '검사받음'이 결과(증상)와 HP 둘 다에 연관되면 collider가 되어 편향.")
CODE("""증상/궤양 ──▶ HP검사받음(selection) ◀── 의사결정
   │
   ▼
 출혈/암      HP상태 ──▶ 출혈/암
# '검사받음'으로 제한(조건화)하면 증상↔HP 가짜경로 개방 가능""")
H("2.3 해결법",2)
NUM("검사 시행률 선확인: pre-KT 정규 EGD·HP 검사율 산출(정규일수록 selection↓).")
NUM("Restriction + 확인자 vs 미확인자 특성 비교표.")
NUM("IPSW(역확률선택가중): 1/P(검사)로 대표성 복원.")
CODE("""ps_test <- glm(tested ~ age+sex+symptom+egd+comorbid, family=binomial, data=all)
all$w_sel <- 1/predict(ps_test, type="response")""")
NUM("다중대치(MI, m=20)로 HP 결측 보정. / 추적손실 IPCW. / new-user로 prevalent 차단.")
H("2.4 예시",2)
P("검사군 출혈률 8% vs 모집단 추정 5%(검사군=증상자 위주)면 8%는 과대. IPSW 가중 후 ~5%로 보정되고 HP 효과추정도 덜 편향.")

H("3. Confounding by Indication (적응증 교란)",1)
P("제균·내시경이 증상/궤양/위축 때문에 시행되면 그 적응증이 노출·결과의 공통원인이 되어 교란.")
NUM("Active comparator: G1·G2 모두 HP 양성 → 'HP 감염' 적응증 공유로 상쇄.")
NUM("성향점수(PS): 제균 확률 모형 → IPTW/매칭. 균형은 SMD<0.1로 확인(p값 대신).")
CODE("""ps <- glm(eradicated ~ age+sex+ulcer+atrophy+symptom+IS_regimen+charlson,
          family=binomial, data=hp_pos)
hp_pos$ptw <- ifelse(hp_pos$eradicated==1, 1/fitted(ps), 1/(1-fitted(ps)))
library(cobalt); bal.tab(eradicated ~ ., data=hp_pos, weights=hp_pos$ptw)
coxph(Surv(t0,t1,bleed) ~ eradicated, data=hp_pos, weights=ptw)""")
NUM("음성대조(negative control): 제균과 무관해야 할 결과/노출에서 효과 보이면 교란 의심.")
NUM("E-value: 관측 HR을 없애는 데 필요한 미측정 교란의 최소 강도. 예 HR=0.6 → E-value≈2.7.")

H("4. Detection / Surveillance Bias",1)
P("이식·제균군이 내시경을 더 자주 받아 무증상 암·출혈을 더 발견 → 노출이 결과를 늘리는(또는 조기발견으로 예후 좋아 보이게) 왜곡.")
BUL("해결: 내시경 횟수 보정(공변량/오프셋), 결과를 임상유의 사건(수혈필요 출혈·침습암)으로 제한, 유사 감시강도 하위집단 분석.")
P("예: 제균군 연 1.5회 vs 미제균 0.7회 내시경. 보정 후 효과 약화/소실 여부로 detection bias 판단.", italic=True, size=9.5)

H("5. Competing Risk Bias (경쟁위험)",1)
P("KT 환자는 사망·이식신장소실이 흔함. 사망 후 암/출혈 불가(경쟁사건) → 표준 1−KM은 누적발생 과대추정. 누적발생함수(CIF)+Fine-Gray(SHR) 사용, cause-specific Cox와 병행.")
CODE("""library(cmprsk)            # status: 0=중도절단,1=위암,2=사망(경쟁)
cuminc(ftime=t, fstatus=status, group=erad)               # CIF
crr(ftime=t, fstatus=status, cov1=X, failcode=1, cencode=0)  # Fine-Gray SHR""")
P("예: 10년 위암 1−KM 4.0%(사망 무시) vs CIF 2.5%(사망 반영). 이식군에서 차이↑ → CIF/Fine-Gray가 정직한 수치.", italic=True, size=9.5)

H("6. Misclassification / Information Bias",1)
BUL("노출 오분류: 비차별적이면 null로 희석(보수적), 차별적이면 방향 불명.")
BUL("결과 오분류: 코드만으로 위암 잡으면 거짓양성↑ → 병리확진+중증등록(V코드) 복합정의로 PPV↑.")
BUL("해결: 검증된 정의(제균 민감도 99.7%, 위암 PPV 94.1%), 정의 바꿔 민감도분석, 음성대조.")

H("7. Lead-time / Length Bias (암 특이)",1)
BUL("Lead-time: 감시 조기진단이 생존을 길어 보이게 → 생존 대신 발생률 비교로 회피.")
BUL("Length bias: 천천히 자라는 암이 감시에 더 포착 → 침습/증상성 결과 제한, 발생률 중심.")

H("8. 통합 분석 워크플로(권고 순서)",1)
for i,s in enumerate([
 "DAG 작성 → 교란/collider/매개 구분, 조정변수(백도어) 결정",
 "Target trial 정의 → 적격·time-zero·전략·결과·추정량 명시",
 "노출 분리: 감염상태(baseline) vs 제균(시간의존)",
 "불멸시간 차단: 시간의존+landmark+CCW 병행",
 "선택 보정: restriction+IPSW/MI(+추적손실 IPCW)",
 "교란 보정: active comparator+PS(IPTW/매칭, SMD<0.1)",
 "경쟁위험: CIF+Fine-Gray",
 "감시/오분류: 내시경빈도 보정·정의 강화",
 "견고성: lag·음성대조·E-value·민감도분석",
 "보고: STROBE/RECORD-PE+흐름도+SMD 균형표"],1):
    NUM(s)

H("9. 한 장 요약표",1)
TABLE(["단계","도구","산출/점검"],
 [["설계","DAG, target trial emulation","조정변수·time-zero·전략"],
  ["불멸시간","tmerge 시간의존 / landmark / CCW(IPCW)","동일결론 수렴 확인"],
  ["선택","restriction + IPSW / MI","확인자 vs 미확인자 표"],
  ["교란","PS-IPTW/매칭, active comparator","SMD<0.1, E-value"],
  ["감시","내시경빈도 보정, 결과강화","보정 전후 HR"],
  ["경쟁위험","CIF, Fine-Gray(SHR)","1−KM과 비교"],
  ["오분류","검증된 정의, 음성대조","민감도분석"]],
 widths=[2.6,6.4,5.5])

H("참고문헌(방법론)",1)
for rr in [
 "Hernán MA, Robins JM. Am J Epidemiol 2016 — Target trial emulation.",
 "Suissa S. Am J Epidemiol 2008 — Immortal time bias.",
 "Hernán MA, et al. Epidemiology 2004 — Selection bias / colliders.",
 "VanderWeele TJ, Ding P. Ann Intern Med 2017 — E-value.",
 "Austin PC, Fine JP. Stat Med 2017 — Competing risks (Fine-Gray).",
 "Maringe C, et al. Int J Epidemiol 2020 — Clone-censor-weight.",
]:
    p=DOC.add_paragraph(style="List Number"); r=p.add_run(rr); _k(r); r.font.size=Pt(9)
P()
P("면책: 교육·설계 보조 문서. 코드는 개념 예시이며 실제 자료에 맞춰 검증·수정 후 사용. 통계전문가 검토 권고.",
  italic=True, size=8.5, color=(0x66,0x66,0x66))

OUT="/home/user/dahyunschedule/Bias_methodology_guide.docx"
DOC.save(OUT); print("Saved:", OUT)
