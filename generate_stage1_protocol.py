#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Stage-1 CMC KT H. pylori protocol as .docx (mirrors the markdown)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()

def set_base_font(doc, name="Malgun Gothic", size=10.5):
    st = doc.styles["Normal"]; st.font.name = name; st.font.size = Pt(size)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts')) or OxmlElement('w:rFonts')
    if rf.getparent() is None: rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), name)
set_base_font(DOC)

def _k(run, name="Malgun Gothic"):
    run.font.name = name
    r = run._element.get_or_add_rPr()
    rf = r.find(qn('w:rFonts')) or OxmlElement('w:rFonts')
    if rf.getparent() is None: r.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), name)

def H(text, level=1):
    h = DOC.add_heading(level=level); _k(h.add_run(text)); return h
def P(text="", bold=False, italic=False, size=None, color=None, align=None):
    p = DOC.add_paragraph(); r = p.add_run(text); _k(r)
    r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor(*color)
    if align: p.alignment=align
    return p
def BUL(text):
    p = DOC.add_paragraph(style="List Bullet"); _k(p.add_run(text)); return p
def NUM(text):
    p = DOC.add_paragraph(style="List Number"); _k(p.add_run(text)); return p
def TABLE(headers, rows, widths=None):
    t = DOC.add_table(rows=1, cols=len(headers)); t.style="Light Grid Accent 1"
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""; r=c.paragraphs[0].add_run(h); _k(r); r.bold=True; r.font.size=Pt(9)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=""; r=cells[i].paragraphs[0].add_run(str(v)); _k(r); r.font.size=Pt(9)
    if widths:
        for i,w in enumerate(widths):
            for row in t.rows: row.cells[i].width=Cm(w)
    DOC.add_paragraph(); return t

# ---- Title ----
t=DOC.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run("서울성모병원 신장이식(KT) 수혜자에서\nHelicobacter pylori 감염·제균과 위장관 출혈 및 위암·대장암 발생"); _k(r); r.bold=True; r.font.size=Pt(16)
s=DOC.add_paragraph(); s.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=s.add_run("단일기관 후향적 코호트 연구 — 1단계 연구계획서(초안)"); _k(r); r.italic=True; r.font.size=Pt(12)
m=DOC.add_paragraph(); m.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=m.add_run("가톨릭대학교 서울성모병원(CMC) · 원내 CDW/EMR · target trial emulation · 작성일 2026-06-10"); _k(r); r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x66,0x66,0x66)
DOC.add_paragraph()

H("1. 연구 배경 및 필요성",1)
BUL("H. pylori 제균은 일반(면역정상) 인구에서 위암을 약 45~55% 감소(RCT·메타분석). 그러나 면역억제 환자에서 제균 여부별 위암·위장관 합병증 차이를 직접 본 연구는 없음(연구 공백).")
BUL("KT 수혜자는 만성 면역억제로 위암 위험이 일반인구 대비 약 1.9~2.3배(한국 신이식 SIR 2.3); 위암은 H. pylori 매개 '감염관련 암'. 이식 후 소화성궤양·위장관 출혈 위험도 높음.")
BUL("H. pylori는 상부위장관 출혈(소화성궤양)의 주원인이며 대장 신생물과의 연관(OR~1.4~1.7)도 보고 → 대장암을 탐색적 결과로 포함.")
BUL("→ 이식 전후 screen-and-treat 임상결정 근거 제공 목적.")

H("2. 연구 목적 및 가설",1)
P("일반목적: 서울성모병원 KT 수혜자에서 H. pylori 감염·제균이 (1) 위장관 출혈, (2) 위암·대장암 발생에 영향을 주는지 검증.")
BUL("Aim 1(출혈): HP+ 미제균군은 제균군·음성군보다 (상부)위장관 출혈 위험↑. H1a: HP+제균<HP+미제균; H1b: HP+미제균>HP−.")
BUL("Aim 2(암): HP+ 미제균군은 제균군보다 위암(탐색적으로 대장암) 발생↑.")
BUL("Aim 3(기술): 위암 조직형(Lauren, EBVaGC, MSI/HER2, PTLD 감별)·대장암 특성 기술.")

H("3. 연구 설계 개요",1)
BUL("단일기관 후향적 코호트. 인과추론 비뚤림(immortal time·selection) 차단 위해 target trial emulation 틀.")
BUL("time-zero(index) = KT 시행일.")
BUL("노출 2축: ① HP 감염상태(baseline), ② 제균치료(이식 전=baseline / 이식 후=시간의존).")
BUL("추적: index → 결과/사망/이식신장소실/추적종료 중 최초.")

H("4. 연구 대상",1)
BUL("세팅: 서울성모병원 성인(≥19세) KT 수혜자. 연구기간(예): 2009–2022 이식, 자료확보 시점까지 추적.")
BUL("포함: 해당 기간 KT 전수. 제외: index 이전 위암/대장암 기왕(해당 암 분석), 위/대장 절제력, HP 상태 완전 미확인(별도군·민감도), 추적 부족, 재이식 중복.")

H("5. 노출 정의",1)
H("5.1 H. pylori 감염상태",2)
BUL("조직/CLO/UBT/대변항원/혈청 IgG 중 1개 이상 양성=HP+, 모두 음성=HP−. 이식 전 정규 EGD 시행률이 확인율·selection을 좌우 → CDW로 선(先)산출.")
H("5.2 제균치료",2)
BUL("PPI/P-CAB+항생제 2종 동시 처방(Park 2023). 성공=치료 후 4주~6개월 확인검사 음성. 시점 층화: 이식 전/후/비제균(이식 후는 시간의존).")
H("5.3 노출군(분석용)",2)
TABLE(["군","정의","주 용도"],
      [["G0","HP 음성","참조(감염효과)"],
       ["G1","HP 양성, 제균(성공)","제균효과 비교"],
       ["G2","HP 양성, 미제균/실패","제균·감염효과 비교"]],
      widths=[2.5,6.5,5.5])
P("주 비교(제균효과): G1 vs G2. 부 비교(감염효과): G2 vs G0.", italic=True, size=9.5)

H("6. 결과 변수",1)
BUL("Aim1 출혈: 상부(주) — 토혈/흑색변 + 내시경 확인 또는 임상적 출혈(Hb≥2g/dL↓·수혈·지혈술·입원); 하부(부) — 혈변+대장내시경. 재출혈·Forrest·수혈량·30일 사망.")
BUL("Aim2 암: 위암(C16, 병리확진; MALT/PTLD 분리), 대장암(C18–C20, 탐색). look-back으로 기왕 제외.")
BUL("Aim3 조직형: 위암 Lauren/WHO/분화도/위치/병기/EBER-ISH/MSI·HER2; 대장암 위치/분화/병기/MSI.")

H("7. 공변량",1)
P("연령·성·이식연도·원인신질환·투석기간·공여자유형·HLA·유도요법(ATG/basiliximab)·유지 면역억제(CNI·스테로이드 누적[프레드니솔론 환산]·MMF/azathioprine·mTOR)·거부반응·CMV/EBV·기저 위병리(위축/장상피화생/궤양)·흡연·음주·BMI·당뇨·암 가족력·PPI/NSAID/aspirin·내시경 감시빈도·eGFR·Charlson.")

H("8. 비뚤림 통제 방법론 (핵심)",1)
H("8.1 Immortal time bias",2)
NUM("시간의존 노출: 이식 후 제균은 처방일 이후만 노출(이전 person-time 비노출), 시간의존 Cox.")
NUM("Landmark 분석: 고정 시점(6/12개월)의 제균 여부로 군 정하고 landmark 이후 추적, 이전 사건·사망 제외.")
NUM("Target trial emulation + clone-censor-weight: 유예기간(grace 6~12개월) 내 제균 시작 vs 미시작 전략으로 복제→어긋나면 인위적 중도절단→IPCW 보정. immortal time 구조적 제거.")
NUM("주의: HP 감염상태(G0 vs G2)는 baseline이라 immortal time 무관 — 시간의존 처리는 제균(G1 vs G2)에 한정.")
H("8.2 Selection bias (가장 중요)",2)
NUM("검사 시행률 선확인: CDW로 pre-KT EGD·HP 검사 시행률 산출(1차 산출물). 정규 EGD면 selection 작음(가설).")
NUM("분석 코호트=HP 확인자, 확인자 vs 미확인자 특성 비교표로 selection 평가.")
NUM("검사 선택 역확률가중(IPSW) 또는 HP 결측 다중대치(MI)로 대표성 보정(민감도).")
H("8.3 Confounding by indication",2)
BUL("active-comparator(G1·G2 모두 HP+), 신규사용자 관점, 성향점수 IPTW/매칭, 기저 위병리·증상 보정.")
H("8.4 Detection/surveillance bias",2)
BUL("내시경 감시빈도 보정, 임상적 유의 사건(수혈필요 출혈·침습암)으로 제한 민감도.")
H("8.5 Competing risk & informative censoring",2)
BUL("사망·이식신장소실 경쟁위험 Fine-Gray 병행; 추적종료 정보성 중도절단 IPCW 민감도.")
H("8.6 기타",2)
BUL("prevalent 암 제외(look-back), 암 lag 0–1년 제외, 제균 성공 한정 민감도, E-value(미측정 교란).")

H("9. 통계분석",1)
BUL("기술: 군별 특성표+SMD, 발생률(per 1,000 PY).")
BUL("Aim1: 시간의존 Cox(제균)·baseline Cox(감염)→HR; Fine-Gray; 재출혈 Andersen-Gill; PS-IPTW; landmark·CCW 병행.")
BUL("Aim2: 시간의존 Cox+Fine-Gray; 사건 적으면 Firth/정확법·발생률·SIR(일반인구 표준화); 대장암 탐색.")
BUL("Aim3: 조직형 빈도·비율(EBVaGC%, Lauren), Fisher.")
BUL("민감도: landmark 6/12개월, CCW grace 6/12, IPSW/MI, 제균성공 한정, lag, 내시경빈도, E-value. R(survival/cmprsk/WeightIt/mice) 또는 SAS.")

H("10. 표본수·검정력·실현가능성",1)
BUL("서울성모병원 KT 누적 ~2,500–3,500(연 ~150–250), HP+ ~40–50%.")
BUL("Aim1 출혈: 사건 비교적 흔함 → 단일기관 비교분석 검정력 가능(주력 종점).")
BUL("Aim2 위암/대장암: 사건 적음 → 발생률·SIR·조직형 기술 1차, 비교 HR은 2단계(NHIS/다기관)에서 확정.")
BUL("사전 검정력: CDW 탐색으로 실제 사건수·노출분율 확정 후 시뮬레이션.")

H("11. 데이터 수집 및 변수(요약 CRF)",1)
TABLE(["도메인","변수","소스"],
 [["식별/인구","가명화ID·연령·성","EMR"],
  ["이식","KT일(index)·공여자·HLA·면역억제·거부반응","이식DB/EMR"],
  ["HP","검사종류·일자·결과·제균처방·확인검사","EMR/처방/검사"],
  ["출혈","증상·내시경소견·Forrest·Hb·수혈·입원·재출혈","EMR/내시경/검사"],
  ["암","위/대장암 진단일·병리(조직형·병기·EBER·MSI·HER2)","병리/암등록"],
  ["공변량","흡연·음주·BMI·당뇨·가족력·동반질환·약물·내시경횟수·eGFR·CMV/EBV","EMR/처방/검사"]],
 widths=[2.8,8.2,3.5])

H("12. 연구 일정(개략)",1)
NUM("0–1개월: 프로토콜 확정·CDW 추출명세·IRB 신청(서울성모병원).")
NUM("2–3개월: CDW 1차 추출, HP/EGD 시행률·사건수 feasibility → 검정력·대조군 전략 확정.")
NUM("4–6개월: 변수 정제·조직형 차트리뷰·분석.")
NUM("7–9개월: 해석·초고(STROBE/RECORD)·CMC 서식 전환 결정.")

H("13. 윤리적 고려",1)
BUL("후향 의무기록/CDW → 서울성모병원 IRB 심의·동의면제 신청(비식별·후향). 가명화·원내 분석. STROBE/RECORD-PE 준수.")

H("14. 한계",1)
BUL("단일기관 외부타당도, HP 검사 selection(완화책 §8.2), confounding by indication, 암 희소 검정력, EBVaGC 판정 가용성, 추적 누락 → 2단계 보완.")

H("15. 2단계(빅데이터) 연계 계획",1)
BUL("1단계 phenotype을 NHIS 청구·검진 연계 또는 다기관 CDM(OMOP)으로 확장해 위암·대장암 비교 HR 확정. 1단계 사건율·효과크기를 2단계 검정력 산정에 활용.")

H("주요 참고문헌",1)
for rr in [
 "Lee YC, et al. Gastroenterology 2016 (제균-위암 메타분석). PMID 26836587",
 "Choi IJ, et al. NEJM 2018/2020 (한국 제균 RCT). PMID 29562147 / 31995688",
 "Jeong S, et al. Sci Rep 2020 (한국 신이식 SIR 2.3). PMC7722878",
 "Engels EA, et al. JAMA 2011 (이식 위암 SIR 1.67). PMID 22045767",
 "Park CH, et al. J Korean Med Sci 2023;38:e278 (제균 조작적 정의). PMC10477078",
 "Yang MS, et al. Cancer Res Treat 2022 (위암 정의검증). PMC9016317",
 "Hernán MA, Robins JM. Am J Epidemiol 2016 (target trial emulation).",
]:
    p=DOC.add_paragraph(style="List Number"); r=p.add_run(rr); _k(r); r.font.size=Pt(9)

P()
P("면책: 연구 기획 보조 초안. 실제 CDW 탐색(사건수·검사 시행률) 후 대조군 전략·검정력·세부 정의를 확정하고 통계·이식·소화기 전문가 및 IRB 검토를 거칠 것.",
  italic=True, size=8.5, color=(0x66,0x66,0x66))

OUT="/home/user/dahyunschedule/Stage1_CMC_KT_Hpylori_protocol.docx"
DOC.save(OUT); print("Saved:", OUT)
